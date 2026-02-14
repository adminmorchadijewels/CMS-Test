import os
from flask import Flask, jsonify, request
from flask_cors import CORS
from openpyxl import Workbook, load_workbook

app = Flask(__name__)
CORS(app)

EXCEL_FILE = os.path.join(os.path.dirname(__file__), "inventory.xlsx")


def ensure_file_exists():
    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()
        ws = wb.active
        ws.append(["Product", "Quantity", "Price"])
        ws.append(["Widget A", 100, 9.99])
        ws.append(["Widget B", 250, 14.50])
        ws.append(["Widget C", 75, 22.00])
        wb.save(EXCEL_FILE)


def read_excel():
    ensure_file_exists()
    wb = load_workbook(EXCEL_FILE)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    columns = [str(c) if c is not None else "" for c in rows[0]]
    data = []
    for row in rows[1:]:
        data.append([v if v is not None else "" for v in row])
    return columns, data


@app.route("/api/data", methods=["GET"])
def get_data():
    columns, data = read_excel()
    return jsonify({"columns": columns, "data": data})


@app.route("/api/data", methods=["PUT"])
def save_data():
    body = request.get_json()
    columns = body["columns"]
    data = body["data"]

    wb = Workbook()
    ws = wb.active
    ws.append(columns)
    for row in data:
        # Pad or trim row to match column count
        padded = list(row) + [""] * (len(columns) - len(row))
        ws.append(padded[: len(columns)])
    wb.save(EXCEL_FILE)
    return jsonify({"ok": True})


@app.route("/api/column", methods=["POST"])
def add_column():
    body = request.get_json()
    col_name = body.get("name", "New Column")

    columns, data = read_excel()
    columns.append(col_name)
    for row in data:
        row.append("")

    wb = Workbook()
    ws = wb.active
    ws.append(columns)
    for row in data:
        ws.append(row)
    wb.save(EXCEL_FILE)
    return jsonify({"columns": columns, "data": data})


@app.route("/api/column/<int:col_index>", methods=["DELETE"])
def delete_column(col_index):
    columns, data = read_excel()
    if col_index < 0 or col_index >= len(columns):
        return jsonify({"error": "Invalid column index"}), 400

    columns.pop(col_index)
    for row in data:
        if col_index < len(row):
            row.pop(col_index)

    wb = Workbook()
    ws = wb.active
    ws.append(columns)
    for row in data:
        ws.append(row)
    wb.save(EXCEL_FILE)
    return jsonify({"columns": columns, "data": data})


if __name__ == "__main__":
    ensure_file_exists()
    app.run(debug=True, port=5000)
